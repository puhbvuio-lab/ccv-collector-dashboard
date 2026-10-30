#!/usr/bin/env python3
"""Twitch CCV 采集守护 + 可视化看板。

功能:
  - 任务管理：从看板添加 Twitch 游戏/频道任务、暂停、恢复、停止
  - 每个任务使用独立线程和独立数据目录 data/<task_id>/
  - 可视化看板 http://127.0.0.1:8088/

用法:
  python server.py
"""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from logging.handlers import TimedRotatingFileHandler
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urlparse

import monitor as monitor_mod
from monitor import (
    BASE_DIR,
    Monitor,
    TwitchClient,
    utcnow,
)

REPORT_THREADS: list[threading.Thread] = []

WEB_NODE_EXE = os.environ.get("CODEX_NODE_EXE") or shutil.which("node") or "node"
WEB_NODE_MODULES = os.environ.get("CODEX_NODE_MODULES", "")


def load_env() -> dict[str, str]:
    env: dict[str, str] = {}
    env_path = BASE_DIR / ".env"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                key, _, value = line.partition("=")
                env[key.strip()] = value.strip().strip('"').strip("'")
    # 进程环境优先于 .env，方便 CI、计划任务和容器化运行。
    for key in ("TWITCH_CLIENT_ID", "TWITCH_CLIENT_SECRET"):
        if os.environ.get(key):
            env[key] = os.environ[key]
    return env


def node_env() -> dict[str, str]:
    """返回网页采集进程环境；优先使用显式 NODE_PATH，否则使用项目本地模块。"""
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    if WEB_NODE_MODULES:
        env["NODE_PATH"] = WEB_NODE_MODULES
    elif not env.get("NODE_PATH"):
        local_modules = BASE_DIR / "node_modules"
        if local_modules.exists():
            env["NODE_PATH"] = str(local_modules)
    return env


def load_config() -> dict[str, Any]:
    cfg_path = BASE_DIR / "config.json"
    return json.loads(cfg_path.read_text(encoding="utf-8")) if cfg_path.exists() else {}


def normalize_external_tool_url(value: object) -> str:
    url = str(value or "").strip()
    if not url or any(char.isspace() for char in url):
        return ""
    try:
        parsed = urlparse(url)
    except ValueError:
        return ""
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        return ""
    if parsed.username is not None or parsed.password is not None:
        return ""
    return url


def slugify(value: str) -> str:
    text = re.sub(r"[^a-z0-9]+", "_", str(value or "").lower()).strip("_")
    return text[:48]


def parse_twitch_channel(target: str) -> str:
    value = str(target or "").strip().lower()
    if "twitch.tv/" in value:
        value = value.split("twitch.tv/", 1)[1]
    value = value.split("/")[0].split("?")[0].split("#")[0].strip("@")
    if not re.fullmatch(r"[a-z0-9_]{1,64}", value):
        raise ValueError(f"无效的 Twitch 频道: {target}")
    return value



# ============================================================ 任务管理


class TaskManager:
    def __init__(self, twitch_creds: dict[str, str]):
        self.tasks_path = BASE_DIR / "data" / "tasks.json"
        self.tasks: dict[str, dict[str, Any]] = {}
        self.threads: dict[str, threading.Thread] = {}
        self.lock = threading.Lock()
        self.twitch_creds = twitch_creds
        self.load()

    def load(self) -> None:
        if self.tasks_path.exists():
            try:
                data = json.loads(self.tasks_path.read_text(encoding="utf-8"))
                self.tasks = {t["id"]: t for t in data.get("tasks", [])}
            except Exception:
                self.tasks = {}
        if not self.tasks:
            default_id = "twitch_ww_default"
            self.tasks[default_id] = {
                "id": default_id,
                "name": "鸣潮Twitch(默认)",
                "platform": "twitch",
                "mode": "game",
                "game_name": "Wuthering Waves",
                "game_id": 209738946,
                "channels": [],
                "max_streams": 3,
                "poll_interval_seconds": 30,
                "monitor_duration_minutes": 0,
                "status": "running",
                "created_at": utcnow().isoformat(),
            }
            self.save()
        changed = False
        for task in list(self.tasks.values()):
            if task.get("platform") != "twitch":
                task["status"] = "unsupported"
                task["status_reason"] = "当前版本已移除内置 YouTube 采集器"
                changed = True
                continue
            if task.get("status") in ("running", "recovering"):
                task["status"] = "recovering"
                self.save()
                self.start_thread(task["id"])
        if changed:
            self.save()

    def save(self) -> None:
        self.tasks_path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.tasks_path.with_suffix(".tmp")
        tmp.write_text(
            json.dumps({"tasks": list(self.tasks.values())}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        os.replace(tmp, self.tasks_path)

    def task_dir(self, task_id: str) -> Path:
        return BASE_DIR / "data" / task_id

    def allocate_task_id(self, base_id: str) -> str:
        """Return an unused task id. Caller must hold ``self.lock``."""
        if base_id not in self.tasks:
            return base_id
        suffix = 2
        while f"{base_id}_{suffix}" in self.tasks:
            suffix += 1
        return f"{base_id}_{suffix}"

    def create(self, payload: dict[str, Any]) -> dict[str, Any]:
        platform = str(payload.get("platform", "twitch")).lower()
        mode = str(payload.get("mode", "")).lower()
        target = str(payload.get("target", "")).strip()
        name = str(payload.get("name", "")).strip() or target
        if platform != "twitch":
            raise ValueError("当前项目只支持 Twitch 采集")
        legacy_interval = max(5, int(payload.get("poll_interval_seconds", 30) or 30))
        watch_interval = max(5, int(payload.get("watch_interval_seconds", legacy_interval) or legacy_interval))
        live_interval = max(5, int(payload.get("live_interval_seconds", legacy_interval) or legacy_interval))
        task: dict[str, Any] = {
            "platform": platform,
            "mode": mode,
            "name": name,
            "max_streams": max(1, int(payload.get("max_streams", 3) or 3)),
            "watch_interval_seconds": watch_interval,
            "live_interval_seconds": live_interval,
            "poll_interval_seconds": live_interval,
            "monitor_duration_minutes": max(0, int(payload.get("monitor_duration_minutes", 0) or 0)),
            "status": "running",
            "created_at": utcnow().isoformat(),
        }
        if mode == "channel":
            task["channels"] = [parse_twitch_channel(target)]
            task["target"] = target
            task_id = f"twitch_{slugify(task['channels'][0])}_{datetime.now().strftime('%H%M%S')}"
        else:
            task["mode"] = "game"
            task["game_name"] = target or "Wuthering Waves"
            task["target"] = task["game_name"]
            task_id = f"twitch_game_{slugify(task['game_name'])}_{datetime.now().strftime('%H%M%S')}"
        if payload.get("dual"):
            if task.get("mode") == "game":
                raise ValueError("dual 模式仅支持指定频道/视频，请使用直播间链接")
            task["dual"] = True
        with self.lock:
            task_id = self.allocate_task_id(task_id)
            task["id"] = task_id
            self.tasks[task_id] = task
            self.save()
            self.start_thread(task_id)
        return task

    def start_thread(self, task_id: str) -> None:
        thread = threading.Thread(target=self._worker, args=(task_id,), daemon=True, name=f"task-{task_id}")
        self.threads[task_id] = thread
        task = self.tasks.get(task_id)
        if task:
            task["status"] = "running"
            self.save()
        thread.start()

    def _worker(self, task_id: str) -> None:
        task = self.tasks.get(task_id)
        if not task:
            return
        data_dir = f"data/{task_id}"
        try:
            if task.get("dual"):
                self._run_dual(task, data_dir)
            else:
                Monitor(self._api_cfg(task, data_dir), self._api_client(task, data_dir)).run()
        except Exception as exc:  # noqa: BLE001
            logging.exception("[%s] 任务异常退出: %s", task_id, exc)
        finally:
            with self.lock:
                task["status"] = "stopped"
                task["stopped_at"] = utcnow().isoformat()
                self.save()
                self.threads.pop(task_id, None)
            logging.info("[%s] 任务线程结束", task_id)
            if task.get("_delete_pending"):
                return
            if load_config().get("auto_report_on_stop", True) and self.task_has_data(task_id):
                thread = threading.Thread(
                    target=self._run_report,
                    args=(task_id,),
                    daemon=True,
                    name=f"report-{task_id}",
                )
                REPORT_THREADS.append(thread)
                thread.start()

    def _api_client(self, task: dict[str, Any], data_dir: str) -> TwitchClient:
        return TwitchClient(
            self.twitch_creds["client_id"],
            self.twitch_creds["client_secret"],
            BASE_DIR / data_dir / "twitch_token.json",
        )

    def _api_cfg(self, task: dict[str, Any], data_dir: str) -> dict[str, Any]:
        legacy_interval = max(5, int(task.get("poll_interval_seconds", 30) or 30))
        watch_interval = max(5, int(task.get("watch_interval_seconds", legacy_interval) or legacy_interval))
        live_interval = max(5, int(task.get("live_interval_seconds", legacy_interval) or legacy_interval))
        if task.get("dual"):
            watch_interval = live_interval
        return {
            "task_id": task["id"],
            "data_dir": data_dir,
            "charts_dir": f"charts/{task['id']}",
            "game_name": task.get("game_name", "Wuthering Waves"),
            "game_id": task.get("game_id", 209738946),
            "languages": [],
            "channels": task.get("channels", []),
            "max_streams": task.get("max_streams", 3),
            "watch_interval_seconds": watch_interval,
            "live_interval_seconds": live_interval,
            "poll_interval_seconds": live_interval,
            "monitor_duration_minutes": task.get("monitor_duration_minutes", 0),
            "auto_charts_on_exit": False,
        }

    def _api_worker(self, task: dict[str, Any], cfg: dict[str, Any]) -> None:
        Monitor(cfg, self._api_client(task, cfg["data_dir"])).run()

    def _run_dual(self, task: dict[str, Any], data_dir: str) -> None:
        task_dir = BASE_DIR / data_dir
        task_dir.mkdir(parents=True, exist_ok=True)
        interval = int(task.get("poll_interval_seconds", 30) or 30)
        stop_file = str(task_dir / "stop.flag")
        pause_file = str(task_dir / "pause.flag")
        cfg = self._api_cfg(task, data_dir)
        api_thread = threading.Thread(
            target=self._api_worker,
            args=(task, cfg),
            daemon=True,
            name=f"api-{task['id']}",
        )
        api_thread.start()
        script = str(BASE_DIR / "twitch_web_pilot.js")
        target = (task.get("channels") or [""])[0] or task.get("target", "")
        web_out = str(task_dir / "web")
        env = node_env()
        logging.info("[%s] dual 启动: API 线程 + Web(%s)", task["id"], script)
        proc = subprocess.Popen(
            [
                WEB_NODE_EXE,
                script,
                target,
                str(interval),
                "0",
                web_out,
                "--stop-file",
                stop_file,
                "--pause-file",
                pause_file,
            ],
            cwd=BASE_DIR,
            env=env,
        )
        api_thread.join()
        try:
            proc.wait(timeout=60)
        except Exception:
            try:
                proc.kill()
            except Exception:
                pass

    def task_has_data(self, task_id: str) -> bool:
        task_dir = self.task_dir(task_id)
        summary = task_dir / "streams_summary.csv"
        if summary.exists():
            try:
                with summary.open("r", encoding="utf-8-sig", newline="") as fh:
                    if sum(1 for _ in csv.DictReader(fh)) > 0:
                        return True
            except Exception:
                pass
        web_csv = task_dir / "web" / "streams.csv"
        if web_csv.exists():
            try:
                with web_csv.open("r", encoding="utf-8") as fh:
                    if sum(1 for _ in fh) > 1:
                        return True
            except Exception:
                pass
        return False

    def _run_report(self, task_id: str) -> None:
        try:
            import subprocess

            env = dict(os.environ)
            env["PYTHONIOENCODING"] = "utf-8"
            result = subprocess.run(
                [sys.executable, str(BASE_DIR / "report.py"), "--task", task_id],
                cwd=BASE_DIR,
                env=env,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=600,
            )
            if result.stdout:
                logging.info("[报告] %s", result.stdout.strip())
            if result.returncode != 0 and result.stderr:
                logging.warning("[报告失败] %s", result.stderr.strip()[-1000:])
        except Exception as exc:  # noqa: BLE001
            logging.exception("[报告失败] %s", exc)

    def control(self, task_id: str, action: str) -> dict[str, Any]:
        task = self.tasks.get(task_id)
        if not task:
            raise KeyError(f"任务不存在: {task_id}")
        if task.get("status") == "unsupported" and action == "resume":
            raise ValueError("该历史任务类型已不支持恢复")
        task_dir = self.task_dir(task_id)
        task_dir.mkdir(parents=True, exist_ok=True)
        if action == "pause":
            (task_dir / "pause.flag").write_text("", encoding="utf-8")
            task["status"] = "paused"
        elif action == "resume":
            flag = task_dir / "pause.flag"
            if flag.exists():
                flag.unlink()
            stop_flag = task_dir / "stop.flag"
            if stop_flag.exists():
                stop_flag.unlink()
            task["status"] = "running"
            if task_id not in self.threads or not self.threads[task_id].is_alive():
                self.start_thread(task_id)
        elif action == "stop":
            (task_dir / "stop.flag").write_text("", encoding="utf-8")
            task["status"] = "stopping"
        else:
            raise ValueError(f"未知操作: {action}")
        self.save()
        return task

    def delete(self, task_id: str, purge: bool = False) -> dict[str, Any]:
        with self.lock:
            task = self.tasks.get(task_id)
            if not task:
                raise KeyError(f"任务不存在: {task_id}")
            worker = self.threads.get(task_id)
            if (worker is not None and worker.is_alive()) or task.get("status") in ("running", "recovering", "paused", "stopping"):
                # 运行中: 先优雅停止，线程结束后由后台线程完成删除
                task["_delete_pending"] = True
                task_dir = self.task_dir(task_id)
                task_dir.mkdir(parents=True, exist_ok=True)
                (task_dir / "stop.flag").write_text("", encoding="utf-8")
                if task.get("status") == "paused":
                    task["status"] = "running"  # 让 run 循环恢复以便感知 stop.flag
                    pause_flag = task_dir / "pause.flag"
                    if pause_flag.exists():
                        pause_flag.unlink()
                else:
                    task["status"] = "stopping"
                self.save()
                threading.Thread(
                    target=self._delete_after_stop,
                    args=(task_id, purge),
                    daemon=True,
                    name=f"delete-{task_id}",
                ).start()
                logging.info("[%s] 删除请求已受理，正在停止任务...", task_id)
                return {"ok": True, "id": task_id, "pending": True}
            self.tasks.pop(task_id, None)
            self.save()
        logging.info("[%s] 任务已删除%s", task_id, "（含采集数据）" if purge else "")
        if purge:
            shutil.rmtree(BASE_DIR / "data" / task_id, ignore_errors=True)
            shutil.rmtree(BASE_DIR / "charts" / task_id, ignore_errors=True)
        return {"ok": True, "id": task_id, "pending": False}

    def _delete_after_stop(self, task_id: str, purge: bool) -> bool:
        thread = self.threads.get(task_id)
        if thread is not None:
            thread.join(timeout=90)
            if thread.is_alive():
                with self.lock:
                    task = self.tasks.get(task_id)
                    if task:
                        task["status"] = "delete_failed"
                        task["delete_error"] = "任务未能在 90 秒内停止，未删除任务或数据"
                        task.pop("_delete_pending", None)
                        self.save()
                logging.error("[%s] 删除超时：任务线程仍在运行，已保留任务和数据", task_id)
                return False
        with self.lock:
            self.tasks.pop(task_id, None)
            self.save()
        if purge:
            shutil.rmtree(BASE_DIR / "data" / task_id, ignore_errors=True)
            shutil.rmtree(BASE_DIR / "charts" / task_id, ignore_errors=True)
        logging.info("[%s] 任务已删除%s", task_id, "（含采集数据）" if purge else "")
        return True

    # ---------- CCV 趋势数据 ----------

    def list_streams(self, task_id: str) -> list[dict[str, Any]]:
        streams_dir = BASE_DIR / "data" / task_id / "streams"
        out: list[dict[str, Any]] = []
        if not streams_dir.exists():
            return out
        for p in sorted(streams_dir.glob("*.csv")):
            stem = p.stem
            sid = stem.rsplit("_", 1)[-1]
            channel = stem[: -(len(sid) + 1)] if len(sid) + 1 < len(stem) else stem
            rows = 0
            first_ts = ""
            last_ts = ""
            try:
                with p.open("r", encoding="utf-8-sig", newline="") as fh:
                    reader = csv.reader(fh)
                    next(reader, None)
                    for row in reader:
                        if not row:
                            continue
                        rows += 1
                        if not first_ts:
                            first_ts = row[0]
                        last_ts = row[0]
            except Exception:
                continue
            if rows:
                out.append(
                    {
                        "stream_id": sid,
                        "channel": channel,
                        "rows": rows,
                        "first_ts": first_ts,
                        "last_ts": last_ts,
                    }
                )
        out.sort(key=lambda x: x["last_ts"], reverse=True)
        return out

    def read_series(
        self,
        task_id: str,
        stream_id: str,
        window_secs: int = 0,
        bucket_secs: int = 0,
    ) -> list[dict[str, Any]]:
        if not stream_id:
            return []
        streams_dir = BASE_DIR / "data" / task_id / "streams"
        path: Path | None = None
        if streams_dir.exists():
            for p in streams_dir.glob(f"*_{stream_id}.csv"):
                path = p
                break
        if not path:
            return []
        pts: list[tuple[datetime, float]] = []
        now = utcnow()
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            for row in csv.DictReader(fh):
                try:
                    ts = datetime.fromisoformat(row["ts_utc"])
                    ccv = float(row.get("viewer_count") or 0)
                except Exception:
                    continue
                if window_secs > 0 and (now - ts).total_seconds() > window_secs:
                    continue
                pts.append((ts, ccv))
        pts.sort(key=lambda x: x[0])
        if bucket_secs > 0:
            agg: dict[int, list[float]] = {}
            for ts, ccv in pts:
                key = int(ts.timestamp()) // bucket_secs
                agg.setdefault(key, []).append(ccv)
            return [
                {
                    "t": datetime.fromtimestamp(k * bucket_secs, tz=timezone.utc).isoformat(),
                    "ccv": round(sum(v) / len(v), 1),
                }
                for k, v in sorted(agg.items())
            ]
        return [{"t": ts.isoformat(), "ccv": ccv} for ts, ccv in pts]

    def export_xlsx(self, task_id: str) -> bytes:
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference, Series
        from openpyxl.chart.data_source import AxDataSource, StrRef
        from openpyxl.styles import Font

        streams = self.list_streams(task_id)

        summary_by_stream: dict[str, dict[str, str]] = {}
        summary_path = self.task_dir(task_id) / "streams_summary.csv"
        if summary_path.exists():
            try:
                with summary_path.open("r", encoding="utf-8-sig", newline="") as fh:
                    summary_by_stream = {
                        str(row.get("stream_id", "")): row for row in csv.DictReader(fh)
                    }
            except Exception:
                summary_by_stream = {}

        def excel_safe_text(value: Any) -> str:
            text = str(value or "")
            return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text

        def beijing_time(value: str) -> datetime | str:
            raw = str(value or "")
            try:
                parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
                if parsed.tzinfo is None:
                    parsed = parsed.replace(tzinfo=timezone.utc)
                return parsed.astimezone(timezone(timedelta(hours=8))).replace(tzinfo=None)
            except (TypeError, ValueError):
                return excel_safe_text(raw)

        wb = Workbook()
        ws = wb.active
        ws.title = "采样数据"
        ws.append(["采样时间（北京时间）", "直播链接", "主播名称", "直播标题", "CCV", "直播场次ID", "时间标签（图表用）"])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        chart_ranges: list[tuple[int, int, str, str]] = []
        for s in streams:
            stream_id = str(s["stream_id"])
            channel = str(s["channel"])
            meta = summary_by_stream.get(stream_id, {})
            display_name = excel_safe_text(meta.get("display_name") or channel)
            fallback_title = meta.get("title", "")
            stream_path = self.task_dir(task_id) / "streams" / f"{channel}_{stream_id}.csv"
            first_row = ws.max_row + 1
            with stream_path.open("r", encoding="utf-8-sig", newline="") as fh:
                for row in csv.DictReader(fh):
                    try:
                        ccv = float(row.get("viewer_count") or "")
                    except (TypeError, ValueError):
                        continue
                    time_cell = beijing_time(row.get("ts_utc", ""))
                    time_label = time_cell.strftime("%m-%d %H:%M:%S") if isinstance(time_cell, datetime) else str(time_cell)
                    ws.append(
                        [
                            time_cell,
                            f"https://www.twitch.tv/{quote(channel, safe='')}",
                            display_name,
                            excel_safe_text(row.get("title") or fallback_title),
                            ccv,
                            excel_safe_text(stream_id),
                        ]
                    )
                    row_number = ws.max_row
                    ws.cell(row_number, 1).number_format = "yyyy-mm-dd hh:mm:ss"
                    ws.cell(row_number, 7).value = time_label
                    link_cell = ws.cell(row_number, 2)
                    link_cell.hyperlink = link_cell.value
                    link_cell.style = "Hyperlink"
            if ws.max_row >= first_row:
                chart_ranges.append((first_row, ws.max_row, display_name, stream_id))

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 48
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 20
        ws.column_dimensions["G"].hidden = True  # 图表用文本标签列，隐藏不碍观感

        chart_ws = wb.create_sheet("CCV趋势图")
        chart_ws["A1"] = "Excel 原生图表：可编辑标题、坐标轴、样式和数据范围"
        chart_ws["A1"].font = Font(bold=True)
        if chart_ranges:
            chart = LineChart()
            chart.title = "CCV 趋势"
            chart.style = 13
            chart.height = 14
            chart.width = 28
            # 用文本分类轴（CatAx）而不是日期轴（dateAx）：WPS/Excel 对 dateAx
            # 的自动缩放处理不一致（按"天"缩放会把数小时的数据压成一条竖线）。
            # 分类轴每点等距、引用隐藏的文本标签列，任何渲染器都稳定。
            chart.x_axis.title = "采样时间（北京时间）"
            first_count = chart_ranges[0][1] - chart_ranges[0][0] + 1
            skip = max(1, round(first_count / 8))
            chart.x_axis.tickLblSkip = skip
            chart.x_axis.tickMarkSkip = skip
            chart.x_axis.delete = False
            chart.y_axis.title = "CCV"
            chart.y_axis.scaling.min = 0
            chart.y_axis.delete = False
            chart.visible_cells_only = False  # 标签列 G 被隐藏，仍需参与绘图
            for first_row, last_row, display_name, stream_id in chart_ranges:
                values = Reference(ws, min_col=5, min_row=first_row, max_row=last_row)
                labels = Reference(ws, min_col=7, min_row=first_row, max_row=last_row)
                series = Series(values, title=f"{display_name} · {stream_id}")
                series.cat = AxDataSource(strRef=StrRef(f=str(labels)))
                chart.series.append(series)
            chart_ws.add_chart(chart, "A3")

        bio = BytesIO()
        wb.save(bio)
        wb.close()
        return bio.getvalue()

    def snapshot(self) -> list[dict[str, Any]]:
        result = []
        for task in self.tasks.values():
            item = dict(task)
            item["thread_alive"] = self.threads.get(task["id"]).is_alive() if self.threads.get(task["id"]) else False
            item["heartbeat"] = self.read_heartbeat(task["id"])
            item["live_streams"] = self.read_live_streams(task["id"])
            result.append(item)
        return result

    def read_heartbeat(self, task_id: str) -> dict[str, Any]:
        path = self.task_dir(task_id) / "heartbeat.json"
        if not path.exists():
            return {}
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}

    def read_live_streams(self, task_id: str) -> list[dict[str, Any]]:
        summary_path = self.task_dir(task_id) / "streams_summary.csv"
        if not summary_path.exists():
            return []
        try:
            with summary_path.open("r", encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.DictReader(fh))
        except Exception:
            return []
        streams = []
        for row in rows:
            if row.get("status") != "live":
                continue
            stream_id = row.get("stream_id", "")
            series = self.read_stream_tail(task_id, stream_id)
            streams.append(
                {
                    "stream_id": stream_id,
                    "channel": row.get("channel", ""),
                    "display_name": row.get("display_name", ""),
                    "current_ccv": series[-1]["ccv"] if series else None,
                    "peak_ccv": row.get("peak_ccv", ""),
                    "avg_ccv": row.get("avg_ccv", ""),
                    "sample_count": row.get("sample_count", ""),
                    "duration_min": row.get("duration_min", ""),
                    "series": series,
                }
            )
        return streams

    def read_stream_tail(self, task_id: str, stream_id: str, max_rows: int = 240) -> list[dict[str, Any]]:
        stream_dir = self.task_dir(task_id) / "streams"
        match = list(stream_dir.glob(f"*_{stream_id}.csv")) if stream_dir.exists() else []
        if not match:
            return []
        try:
            with match[0].open("r", encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.DictReader(fh))
            out = []
            for row in rows[-max_rows:]:
                try:
                    out.append({"t": row.get("ts_utc", ""), "ccv": float(row.get("viewer_count", 0) or 0)})
                except Exception:
                    continue
            return out
        except Exception:
            return []

    def shutdown(self) -> None:
        for task in self.tasks.values():
            if task.get("status") in ("running", "paused", "stopping"):
                (self.task_dir(task["id"]) / "stop.flag").write_text("", encoding="utf-8")
        for thread in list(self.threads.values()):
            thread.join(timeout=30)
        self.save()


# ============================================================ 看板页面

PAGE = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>直播 CCV 追踪器 · 本地工作台</title>
<style>
  :root {
    --bg:#faf6f0; --card:#ffffff; --line:#efe5d8; --line2:#e7dcc9;
    --text:#1b1512; --dim:#9a9086; --dim2:#8a8178;
    --accent:#e8492f; --accent-soft:#fdeee9;
    --ok:#2f9e5d; --ok-soft:#e9f6ee;
    --bad:#d43c1e; --bad-soft:#fdecea;
    --warn:#c47a0a; --warn-soft:#fdf3e3;
  }
  * { box-sizing:border-box; }
  html { -webkit-font-smoothing:antialiased; }
  body { margin:0; background:var(--bg); color:var(--text); font-family:"Microsoft YaHei","PingFang SC",system-ui,sans-serif; }
  .wrap { max-width:1100px; margin:0 auto; padding:40px 28px 70px; }

  .kicker { color:var(--accent); font-size:11px; font-weight:800; letter-spacing:2.5px; }
  .headrow { display:flex; justify-content:space-between; align-items:flex-end; gap:16px; flex-wrap:wrap; }
  h1 { font-size:34px; line-height:1.15; margin:10px 0 8px; font-weight:900; letter-spacing:.5px; }
  h1 .hl { color:var(--accent); }
  .sub { color:var(--dim2); font-size:13px; }
  .sub b { color:var(--dim2); font-weight:700; }
  .servicetag { color:var(--ok); font-size:12px; font-weight:700; white-space:nowrap; padding-bottom:4px; }
  .servicetag .dot { display:inline-block; width:7px; height:7px; border-radius:50%; background:var(--ok); margin-right:6px; vertical-align:1px; }
  .servicetag small { color:var(--dim); font-weight:400; margin-left:8px; }

  .card { background:var(--card); border:1px solid var(--line); border-radius:16px; padding:24px 28px; margin-top:20px; box-shadow:0 1px 3px rgba(90,55,25,.05); }
  .sec { display:flex; justify-content:space-between; align-items:baseline; margin-bottom:18px; gap:12px; flex-wrap:wrap; }
  .sec .t { font-size:16px; font-weight:900; }
  .sec .t .num { color:var(--accent); font-weight:900; font-size:12px; margin-right:10px; letter-spacing:1px; }
  .sec .hint { color:var(--dim); font-size:12px; }

  .frow { display:flex; align-items:center; gap:12px; margin-bottom:12px; flex-wrap:wrap; }
  .flabel { width:96px; flex:none; color:var(--dim2); font-size:13px; }
  .fhelp { color:var(--dim); font-size:12px; margin-top:-4px; margin-bottom:12px; padding-left:108px; }
  input,select { background:#fff; color:var(--text); border:1px solid var(--line2); border-radius:8px; padding:9px 12px; font-size:13px; font-family:inherit; }
  input:focus,select:focus { outline:none; border-color:var(--accent); box-shadow:0 0 0 3px rgba(232,73,47,.09); }
  input[type=number] { width:84px; }
  .tin { flex:1; min-width:260px; }
  .chk { display:inline-flex; align-items:center; gap:6px; font-size:13px; color:#6d645a; cursor:pointer; }
  .chk input { width:auto; accent-color:var(--accent); }

  button { border:0; border-radius:8px; padding:9px 18px; cursor:pointer; font-size:13px; font-weight:700; font-family:inherit; transition:filter .15s; }
  button:disabled { opacity:.4; cursor:not-allowed; }
  .btn-primary { background:linear-gradient(180deg,#f4603a,#e8492f); color:#fff; box-shadow:0 3px 8px rgba(232,73,47,.28); }
  .btn-primary:hover:not(:disabled) { filter:brightness(1.06); }
  .btn-ghost { background:#fff; color:#6d645a; border:1px solid var(--line2); }
  .btn-ghost:hover:not(:disabled) { background:#faf5ee; }
  .btn-danger { background:#fff; color:var(--bad); border:1px solid #f2cfc7; }
  .btn-danger:hover:not(:disabled) { background:var(--bad-soft); }
  .btn-sm { padding:6px 14px; font-size:12px; }
  .tool-link { display:inline-flex; align-items:center; border-radius:8px; padding:9px 14px; font-size:12px; font-weight:700; text-decoration:none; background:#fff; color:#6d645a; border:1px solid var(--line2); }
  .tool-link.disabled { opacity:.45; cursor:not-allowed; }
  .seg { display:inline-flex; border:1px solid var(--line2); border-radius:8px; overflow:hidden; }
  .seg button { border-radius:0; background:#fff; color:#8a8178; padding:6px 14px; font-size:12px; font-weight:600; }
  .seg button + button { border-left:1px solid var(--line2); }
  .seg button.on { background:var(--accent); color:#fff; }

  .statgrid { display:grid; grid-template-columns:repeat(6,1fr); gap:14px; margin-top:20px; }
  @media (max-width:960px){ .statgrid { grid-template-columns:repeat(3,1fr); } }
  .stat { background:var(--card); border:1px solid var(--line); border-radius:14px; padding:16px 18px; box-shadow:0 1px 3px rgba(90,55,25,.04); }
  .stat .lab { color:var(--dim); font-size:12px; margin-bottom:10px; }
  .stat .val { font-size:26px; font-weight:900; letter-spacing:-.5px; line-height:1; font-variant-numeric:tabular-nums; }
  .stat .val.ok { color:var(--ok); }
  .stat .val.acc { color:var(--accent); }
  .stat .cap { color:var(--dim); font-size:11px; margin-top:9px; }
  .stat .bar { height:3px; background:#f0e8dc; border-radius:2px; margin-top:10px; overflow:hidden; }
  .stat .bar i { display:block; height:100%; background:var(--ok); border-radius:2px; }

  .trendbar { display:flex; align-items:center; gap:10px; flex-wrap:wrap; margin-bottom:16px; }
  .trendbar select { max-width:240px; }
  .trendbar .sp { flex:1; }
  .trendmeta { color:var(--dim); font-size:12px; margin-bottom:10px; }
  #trendChart { width:100%; }
  #trendEmpty { color:var(--dim); font-size:13px; text-align:center; padding:70px 0; }

  .task-head { display:flex; align-items:center; gap:10px; flex-wrap:wrap; }
  .task-head .tname { font-size:15px; font-weight:900; }
  .task-meta { color:var(--dim); font-size:12px; margin:10px 0 12px; display:flex; gap:16px; flex-wrap:wrap; }
  .task-actions { margin:0 0 14px; display:flex; gap:8px; }
  .pill { display:inline-block; padding:3px 11px; border-radius:999px; font-size:12px; font-weight:700; }
  .pill.ok { background:var(--ok-soft); color:var(--ok); }
  .pill.bad { background:var(--bad-soft); color:var(--bad); }
  .pill.warn { background:var(--warn-soft); color:var(--warn); }
  .pill.info { background:#f4efe7; color:var(--dim2); }
  .mono { font-family:Consolas,monospace; font-size:12px; color:var(--dim); }

  table { width:100%; border-collapse:collapse; font-size:13px; }
  th,td { text-align:left; padding:10px 8px; border-bottom:1px solid #f4ede2; }
  th { color:var(--dim); font-weight:600; font-size:12px; }
  tr:last-child td { border-bottom:0; }
  td b { font-weight:800; }
  .spark { width:150px; height:38px; display:block; }

  #refreshAt { color:var(--dim); font-size:12px; }
</style>
</head>
<body>
<div class="wrap">
  <div class="kicker">LIVE ANALYTICS / 本地工作台</div>
  <div class="headrow">
    <div>
      <h1>直播实时 <span class="hl">CCV</span> 追踪器</h1>
      <div class="sub">实时轮询 · 趋势观测 · <b>Twitch</b> 直播数据 · <span id="refreshAt"></span></div>
    </div>
    <div style="display:flex;gap:10px;align-items:center;flex-wrap:wrap">
      <!-- YOUTUBE_TOOL_LINK -->
      <div class="servicetag"><span class="dot"></span>服务运行中<small>数据仅保存在本机 data 目录</small></div>
    </div>
  </div>

  <div class="card">
    <div class="sec">
      <div class="t"><span class="num">01</span>采集设置</div>
      <div class="hint">添加后立即开始轮询，可随时暂停 / 恢复 / 停止</div>
    </div>
    <div class="frow">
      <div class="flabel">平台与模式</div>
      <input id="platform" type="hidden" value="twitch" />
      <select id="mode">
        <option value="game">游戏分类（自动选最高CCV）</option>
        <option value="channel">指定频道/直播间链接</option>
      </select>
      <input id="name" class="tin" style="max-width:220px" placeholder="任务名（可选）" />
    </div>
    <div class="frow">
      <div class="flabel">采集目标</div>
      <input id="target" class="tin" placeholder="游戏名 / 频道或直播间链接" />
    </div>
    <div class="fhelp">示例：游戏分类 Wuthering Waves；频道 https://www.twitch.tv/tenz</div>
    <div class="frow">
      <div class="flabel">采集参数</div>
      <label>最多跟踪 <input id="maxStreams" type="number" value="3" min="1" max="20" /></label>
      <label>时长 <input id="duration" type="number" value="0" min="0" /> 分钟</label>
      <label>未开播检查间隔 <input id="watchInterval" type="number" value="30" min="5" /> 秒</label>
      <label>开播采样间隔 <input id="liveInterval" type="number" value="30" min="5" /> 秒</label>
    </div>
    <div class="frow">
      <div class="flabel">持续守望</div>
      <span class="fhelp" style="padding-left:0;margin:0">频道任务会持续等待开播；下播后返回守望，直到手动停止</span>
      <span style="flex:1"></span>
      <button class="btn-primary" onclick="addTask()">开始采集 ▸</button>
    </div>
  </div>

  <div class="statgrid" id="statGrid"></div>

  <div class="card">
    <div class="sec">
      <div class="t"><span class="num">02</span>CCV 趋势</div>
      <div class="hint">按时间范围与采样间隔查看每场直播的 CCV 曲线</div>
    </div>
    <div class="trendbar">
      <select id="trendTask" onchange="onTrendTask()"><option value="">选择任务…</option></select>
      <select id="trendStream" onchange="loadTrend()"><option value="">选择直播场次…</option></select>
      <div class="seg" id="rangeSeg"></div>
      <select id="trendBucket" onchange="loadTrend()">
        <option value="0">原始采样点</option>
        <option value="60">每 1 分钟取点</option>
        <option value="300">每 5 分钟取点</option>
        <option value="600">每 10 分钟取点</option>
        <option value="1800">每 30 分钟取点</option>
      </select>
      <span class="sp"></span>
      <button class="btn-ghost btn-sm" onclick="loadTrend()">刷新</button>
      <button class="btn-primary btn-sm" onclick="exportTrendCsv()">导出 CSV ↓</button>
    </div>
    <div class="trendmeta" id="trendMeta"></div>
    <div id="trendChart"></div>
    <div id="trendEmpty">选择上方任务与直播场次后查看 CCV 趋势曲线</div>
  </div>

  <div class="card">
    <div class="sec">
      <div class="t"><span class="num">03</span>采集任务</div>
      <div class="hint">已停止的任务可删除；数据默认保留，删除时可选一并清除</div>
    </div>
    <div id="taskList"></div>
  </div>
</div>

<script>
const TZ = "Asia/Singapore";
let trendTask = "", trendStream = "", trendWindow = 0, trendPoints = [];
const RANGES = [["900","15分钟"],["3600","1小时"],["21600","6小时"],["0","全部"]];

function esc(s){ return String(s ?? "").replace(/[&<>"']/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c])); }
function fmt(iso){ if(!iso) return "-"; return new Date(iso).toLocaleString("zh-CN",{timeZone:TZ,hour12:false}); }
function hm(iso){ if(!iso) return "-"; return new Date(iso).toLocaleTimeString("zh-CN",{timeZone:TZ,hour12:false,hour:"2-digit",minute:"2-digit"}); }
function age(sec){ if(sec==null) return "-"; if(sec<60) return Math.round(sec)+"秒前"; if(sec<3600) return Math.round(sec/60)+"分钟前"; return Math.round(sec/3600)+"小时前"; }
function pill(status){
  if(status==="running") return '<span class="pill ok">运行中</span>';
  if(status==="paused") return '<span class="pill warn">已暂停</span>';
  if(status==="stopping") return '<span class="pill warn">停止中</span>';
  return '<span class="pill bad">已停止</span>';
}
function spark(series){
  if(!series || series.length < 2) return '<span class="mono">数据不足</span>';
  const w=150,h=38,pad=3,vals=series.map(p=>p.ccv),max=Math.max(...vals,1),min=Math.min(...vals,0),span=Math.max(max-min,1);
  const step=(w-pad*2)/(vals.length-1);
  const pts=vals.map((v,i)=>(pad+i*step).toFixed(1)+","+(h-pad-((v-min)/span)*(h-pad*2)).toFixed(1)).join(" ");
  return '<svg class="spark" viewBox="0 0 '+w+' '+h+'" preserveAspectRatio="none">'
    +'<polyline points="'+pad+','+(h-pad)+' '+pts+' '+pad+','+(h-pad)+'" fill="rgba(232,73,47,.07)" stroke="none"/>'
    +'<polyline points="'+pts+'" fill="none" stroke="#e8492f" stroke-width="1.8" stroke-linejoin="round" stroke-linecap="round"/></svg>';
}
function statCard(lab, val, cls, cap, barPct){
  const bar = barPct!=null ? '<div class="bar"><i style="width:'+Math.min(100,barPct)+'%"></i></div>' : '';
  return '<div class="stat"><div class="lab">'+lab+'</div><div class="val '+(cls||"")+'">'+val+'</div>'+(cap?'<div class="cap">'+cap+'</div>':"")+bar+'</div>';
}
async function api(path, opts){
  const r = await fetch(path, opts);
  const data = await r.json().catch(()=>({}));
  if(!r.ok) throw new Error(data.error || ("HTTP "+r.status));
  return data;
}
function renderStats(tasks){
  const running = tasks.filter(t=>t.status==="running").length;
  const live = tasks.reduce((a,t)=>a+((t.heartbeat||{}).live_count ?? (t.live_streams||[]).length),0);
  const cur = tasks.reduce((a,t)=>a+(t.live_streams||[]).reduce((b,s)=>b+(Number(s.current_ccv)||0),0),0);
  const tracked = tasks.reduce((a,t)=>a+((t.heartbeat||{}).tracked_total ?? 0),0);
  const samples = tasks.reduce((a,t)=>a+(t.live_streams||[]).reduce((b,s)=>b+(Number(s.sample_count)||0),0),0);
  const g = document.getElementById("statGrid");
  if(!tasks.length){ g.innerHTML = statCard("采集任务","0","","尚无任务，先在上方添加"); return; }
  g.innerHTML =
    statCard("采集任务", tasks.length, "", "运行中 "+running+" 个") +
    statCard("在播直播间", live, "ok", "实时跟踪中") +
    statCard("当前CCV合计", cur.toLocaleString("zh-CN"), "acc", "所有在播场次求和") +
    statCard("累计跟踪场次", tracked, "", "含已结束直播") +
    statCard("采样点", samples.toLocaleString("zh-CN"), "", "已写入的CCV快照") +
    statCard("运行占比", tasks.length?Math.round(running/tasks.length*100)+"%":"0%", "ok", "运行中 / 任务总数", tasks.length?running/tasks.length*100:0);
}

/* ---------- 趋势图 ---------- */
function renderRangeSeg(){
  document.getElementById("rangeSeg").innerHTML = RANGES.map(([v,l]) =>
    '<button class="'+(Number(v)===trendWindow?"on":"")+'" onclick="setRange('+v+')">'+l+'</button>').join("");
}
function setRange(v){ trendWindow = v; renderRangeSeg(); loadTrend(); }
async function refreshTrendTasks(tasks){
  const sel = document.getElementById("trendTask");
  const cur = trendTask;
  sel.innerHTML = '<option value="">选择任务…</option>' + tasks.map(t=>'<option value="'+esc(t.id)+'">'+esc(t.name)+'</option>').join("");
  if(cur && tasks.some(t=>t.id===cur)) sel.value = cur;
  else if(tasks.length){ trendTask = tasks[0].id; sel.value = trendTask; onTrendTask(); }
}
async function onTrendTask(){
  trendTask = document.getElementById("trendTask").value;
  const sel = document.getElementById("trendStream");
  sel.innerHTML = '<option value="">选择直播场次…</option>';
  trendStream = ""; trendPoints = []; drawTrend();
  if(!trendTask) return;
  try{
    const d = await api("/api/tasks/"+encodeURIComponent(trendTask)+"/streams");
    if(d.streams.length){
      sel.innerHTML = '<option value="">选择直播场次…</option>' + d.streams.map(s=>
        '<option value="'+esc(s.stream_id)+'">'+esc(s.channel)+' · '+esc(s.stream_id)+' · '+s.rows+'点</option>').join("");
      trendStream = d.streams[0].stream_id;
      sel.value = trendStream;
    }
    await loadTrend();
  }catch(e){ document.getElementById("trendMeta").textContent = "加载场次失败: "+e.message; }
}
async function loadTrend(){
  trendStream = document.getElementById("trendStream").value;
  trendTask = document.getElementById("trendTask").value;
  if(!trendTask || !trendStream){ trendPoints=[]; drawTrend(); return; }
  const bucket = Number(document.getElementById("trendBucket").value) || 0;
  try{
    const d = await api("/api/tasks/"+encodeURIComponent(trendTask)+"/series?stream="+encodeURIComponent(trendStream)
      +"&window="+trendWindow+"&bucket="+bucket);
    trendPoints = d.points || [];
    drawTrend();
  }catch(e){ document.getElementById("trendMeta").textContent = "加载趋势失败: "+e.message; }
}
function drawTrend(){
  const chart = document.getElementById("trendChart");
  const empty = document.getElementById("trendEmpty");
  const meta = document.getElementById("trendMeta");
  if(!trendPoints.length){
    chart.innerHTML = ""; empty.style.display = "block";
    meta.textContent = trendTask && trendStream ? "该范围内暂无数据点" : "";
    return;
  }
  empty.style.display = "none";
  const W=880,H=300,pl=64,pr=16,pt=14,pb=34;
  const vals = trendPoints.map(p=>p.ccv);
  const vmax = Math.max(...vals,1), vmin = Math.min(...vals,0), vspan = Math.max(vmax-vmin,1);
  const n = trendPoints.length;
  const t0 = new Date(trendPoints[0].t).getTime(), t1 = new Date(trendPoints[n-1].t).getTime();
  const tspan = Math.max(t1-t0,1);
  const X = i => pl + (i/(Math.max(n-1,1)))*(W-pl-pr);
  const Y = v => pt + (1-(v-vmin)/vspan)*(H-pt-pb);
  let grid="", ylab="";
  for(let i=0;i<=4;i++){
    const v = vmin + vspan*i/4, y = Y(v);
    grid += '<line x1="'+pl+'" y1="'+y.toFixed(1)+'" x2="'+(W-pr)+'" y2="'+y.toFixed(1)+'" stroke="#f0e8dc" stroke-width="1"/>';
    ylab += '<text x="'+(pl-10)+'" y="'+(y+4).toFixed(1)+'" text-anchor="end" font-size="11" fill="#9a9086">'+Math.round(v).toLocaleString("zh-CN")+'</text>';
  }
  const line = trendPoints.map((p,i)=>X(i).toFixed(1)+","+Y(p.ccv).toFixed(1)).join(" ");
  const area = pl+","+(H-pb)+" "+line+" "+X(n-1).toFixed(1)+","+(H-pb);
  const dots = n<=60 ? trendPoints.map((p,i)=>'<circle cx="'+X(i).toFixed(1)+'" cy="'+Y(p.ccv).toFixed(1)+'" r="2.2" fill="#e8492f"/>').join("") : "";
  let xlab="";
  for(let i=0;i<4;i++){
    const idx = Math.round(i*(n-1)/3);
    if(n<2 || (i>0 && idx===Math.round((i-1)*(n-1)/3))) continue;
    xlab += '<text x="'+X(idx).toFixed(1)+'" y="'+(H-10)+'" text-anchor="middle" font-size="11" fill="#9a9086">'+hm(trendPoints[idx].t)+'</text>';
  }
  chart.innerHTML = '<svg viewBox="0 0 '+W+' '+H+'" preserveAspectRatio="xMidYMid meet">'
    + grid + ylab + xlab
    + '<polyline points="'+area+'" fill="rgba(232,73,47,.08)" stroke="none"/>'
    + '<polyline points="'+line+'" fill="none" stroke="#e8492f" stroke-width="2" stroke-linejoin="round" stroke-linecap="round"/>'
    + dots + '</svg>';
  meta.textContent = "场次 "+trendStream+" · "+n+" 个数据点 · "+fmt(trendPoints[0].t)+" ~ "+fmt(trendPoints[n-1].t)
    +" · 峰值 "+Math.max(...vals).toLocaleString("zh-CN");
}
function exportTrendCsv(){
  if(!trendPoints.length){ alert("当前趋势没有可导出的数据"); return; }
  const lines = ["ts_utc,ccv"].concat(trendPoints.map(p=>p.t+","+p.ccv));
  const blob = new Blob(["\\ufeff"+lines.join("\\n")], {type:"text/csv;charset=utf-8"});
  const a = document.createElement("a");
  a.href = URL.createObjectURL(blob);
  a.download = "ccv_trend_"+trendTask+"_"+trendStream+".csv";
  a.click();
  URL.revokeObjectURL(a.href);
}

/* ---------- 任务列表 ---------- */
async function load(){
  try{
    const data = await api("/api/tasks");
    document.getElementById("refreshAt").textContent = "自动刷新 · 上次 " + new Date().toLocaleTimeString("zh-CN",{hour12:false});
    const tasks = data.tasks || [];
    renderStats(tasks);
    refreshTrendTasks(tasks);
    const list = document.getElementById("taskList");
    list.innerHTML = "";
    if(!tasks.length){ list.innerHTML = '<div class="mono" style="padding:20px 0;text-align:center">暂无任务 · 在上方「采集设置」中添加第一个采集任务</div>'; return; }
    for(const t of tasks){
      const hb = t.heartbeat || {};
      const streams = t.live_streams || [];
      const watching = hb.phase==="waiting" && t.status==="running";
      const div = document.createElement("div");
      div.style.cssText = "border:1px solid #f4ede2;border-radius:12px;padding:18px 20px;margin-bottom:14px;background:#fffdfb";
      div.innerHTML = `
        <div class="task-head">${pill(t.status)}
          ${watching?'<span class="pill warn">守望中 · 等待开播</span>':""}
          <span class="tname">${esc(t.name)}</span>
          <span class="pill info">${t.platform}/${t.mode}</span> <span class="mono">${t.id}</span></div>
        <div class="task-meta">
          <span>PID ${hb.pid || "-"}</span>
          <span>上次轮询 ${age(data.now_ms ? Math.max(0,(Date.now()-new Date(hb.last_tick_utc).getTime())/1000) : null)}</span>
          <span>会话 ${hb.session_id || "-"}</span>
          <span>在播 ${hb.live_count ?? 0} / 累计 ${hb.tracked_total ?? 0}</span>
          <span>守望/采样 ${hb.watch_interval_seconds ?? t.watch_interval_seconds ?? "-"}s / ${hb.live_interval_seconds ?? t.live_interval_seconds ?? "-"}s</span>
          ${hb.remaining_seconds!=null ? `<span>剩余 ${Math.ceil(hb.remaining_seconds/60)} 分钟</span>` : ""}
        </div>
        <div class="task-actions">
          <button class="btn-ghost btn-sm" ${t.status==="running"?"":"disabled"} onclick="ctrl('${t.id}','pause')">暂停</button>
          <button class="btn-primary btn-sm" ${t.status==="paused"?"":"disabled"} onclick="ctrl('${t.id}','resume')">继续</button>
          <button class="btn-danger btn-sm" ${t.status==="stopping"?"disabled":""} onclick="ctrl('${t.id}','stop')">停止</button>
          <button class="btn-danger btn-sm" onclick="delTask('${t.id}')">删除</button>
          <button class="btn-ghost btn-sm" onclick="jumpTrend('${t.id}')">查看趋势</button>
          <button class="btn-primary btn-sm" onclick="exportXlsx('${t.id}')">导出 Excel ↓</button>
        </div>
        <table><thead><tr><th>主播</th><th>当前CCV</th><th>峰值</th><th>均值</th><th>采样</th><th>时长(分)</th><th>曲线</th></tr></thead>
        <tbody>${streams.length===0 ? '<tr><td colspan="7" class="mono">'+(watching?'守望中：正在定时检测开播状态，开播后自动开始采集':'暂无跟踪中的直播')+'</td></tr>' : streams.map(s=>`
          <tr><td><b>${esc(s.channel)}</b></td><td>${s.current_ccv ?? "-"}</td><td>${s.peak_ccv ?? "-"}</td><td>${s.avg_ccv ?? "-"}</td><td>${s.sample_count ?? 0}</td><td>${s.duration_min ?? "-"}</td><td>${spark(s.series)}</td></tr>`).join("")}
        </tbody></table>`;
      list.appendChild(div);
    }
  }catch(e){
    document.getElementById("taskList").innerHTML = '<div class="mono">加载失败: '+esc(e.message)+'</div>';
  }
}
function jumpTrend(id){
  trendTask = id;
  document.getElementById("trendTask").value = id;
  onTrendTask();
  document.getElementById("trendChart").scrollIntoView({behavior:"smooth", block:"center"});
}
function exportXlsx(id){
  const a = document.createElement("a");
  a.href = "/api/tasks/"+encodeURIComponent(id)+"/export.xlsx";
  a.download = "ccv_"+id+".xlsx";
  document.body.appendChild(a);
  a.click();
  a.remove();
}
async function ctrl(id, action){
  try{ await api("/api/tasks/"+encodeURIComponent(id)+"/"+action, {method:"POST"}); load(); }
  catch(e){ alert(e.message); }
}
async function delTask(id){
  if(!confirm("确定删除任务 "+id+" ？\\n若任务正在运行会先自动停止再删除。")) return;
  const purge = confirm("是否同时删除该任务的采集数据目录？\\n\\n确定 = 删除任务并清除数据\\n取消 = 仅从看板移除任务（保留数据）");
  try{
    const r = await api("/api/tasks/"+encodeURIComponent(id)+"?purge="+(purge?1:0), {method:"DELETE"});
    if(trendTask===id){ trendTask=""; trendStream=""; trendPoints=[]; drawTrend(); }
    if(r.pending) alert("任务正在停止，停止后将自动完成删除。");
    load();
  }catch(e){ alert(e.message); }
}
async function addTask(){
  const payload = {
    platform: document.getElementById("platform").value,
    mode: document.getElementById("mode").value,
    target: document.getElementById("target").value.trim(),
    name: document.getElementById("name").value.trim(),
    max_streams: Number(document.getElementById("maxStreams").value),
    monitor_duration_minutes: Number(document.getElementById("duration").value),
    watch_interval_seconds: Number(document.getElementById("watchInterval").value),
    live_interval_seconds: Number(document.getElementById("liveInterval").value),
  };
  if(!payload.target){ alert("请填写目标"); return; }
  try{
    await api("/api/tasks", {method:"POST", headers:{"Content-Type":"application/json"}, body: JSON.stringify(payload)});
    document.getElementById("target").value=""; document.getElementById("name").value="";
    load();
  }catch(e){ alert(e.message); }
}
renderRangeSeg();
load();
setInterval(load, 10000);
setInterval(()=>{ if(trendTask && trendStream) loadTrend(); }, 30000);
</script>
</body>
</html>
"""


def build_index_html(youtube_tool_url: str) -> str:
    url = normalize_external_tool_url(youtube_tool_url)
    if url:
        link = (
            f'<a class="tool-link" href="{html.escape(url, quote=True)}" '
            'target="_blank" rel="noopener noreferrer">打开 YouTube 采集工具 ↗</a>'
        )
    else:
        link = '<span class="tool-link disabled" aria-disabled="true">YouTube 采集工具（待配置）</span>'
    return PAGE.replace("<!-- YOUTUBE_TOOL_LINK -->", link)


class Handler(BaseHTTPRequestHandler):
    manager: TaskManager
    page = PAGE

    def log_message(self, *args):
        pass

    def _json(self, payload: Any, code: int = 200) -> None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _body(self) -> dict[str, Any]:
        length = int(self.headers.get("Content-Length") or 0)
        if length <= 0:
            return {}
        return json.loads(self.rfile.read(length).decode("utf-8"))

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/":
            data = self.page.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        if path == "/api/tasks":
            self._json({"tasks": self.manager.snapshot(), "now_ms": int(time.time() * 1000)})
            return
        match = re.fullmatch(r"/api/tasks/([^/]+)/streams", path)
        if match:
            self._json({"streams": self.manager.list_streams(match.group(1))})
            return
        match = re.fullmatch(r"/api/tasks/([^/]+)/series", path)
        if match:
            qs = parse_qs(parsed.query)
            try:
                window = int(qs.get("window", ["0"])[0] or 0)
                bucket = int(qs.get("bucket", ["0"])[0] or 0)
            except ValueError:
                window, bucket = 0, 0
            stream_id = qs.get("stream", [""])[0]
            points = self.manager.read_series(match.group(1), stream_id, window, bucket)
            self._json({"stream_id": stream_id, "points": points})
            return
        match = re.fullmatch(r"/api/tasks/([^/]+)/export\.xlsx", path)
        if match:
            try:
                data = self.manager.export_xlsx(match.group(1))
            except ValueError as exc:
                self._json({"error": str(exc)}, 400)
                return
            except Exception as exc:  # noqa: BLE001
                self._json({"error": str(exc)}, 500)
                return
            filename = f"ccv_{match.group(1)}.xlsx"
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{quote(filename)}")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        self._json({"error": "not found"}, 404)

    def do_DELETE(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path
        match = re.fullmatch(r"/api/tasks/([^/]+)", path)
        if match:
            purge = parse_qs(parsed.query).get("purge", ["0"])[0] in ("1", "true")
            try:
                self.manager.delete(match.group(1), purge=purge)
                self._json({"ok": True})
            except KeyError as exc:
                self._json({"error": str(exc)}, 404)
            except Exception as exc:  # noqa: BLE001
                self._json({"error": str(exc)}, 400)
            return
        self._json({"error": "not found"}, 404)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/api/tasks":
            try:
                task = self.manager.create(self._body())
                self._json(task, 201)
            except Exception as exc:  # noqa: BLE001
                self._json({"error": str(exc)}, 400)
            return
        match = re.fullmatch(r"/api/tasks/([^/]+)/(pause|resume|stop)", path)
        if match:
            try:
                task = self.manager.control(match.group(1), match.group(2))
                self._json(task)
            except KeyError as exc:
                self._json({"error": str(exc)}, 404)
            except Exception as exc:  # noqa: BLE001
                self._json({"error": str(exc)}, 400)
            return
        self._json({"error": "not found"}, 404)


def main() -> None:
    cfg = load_config()
    parser = argparse.ArgumentParser(description="鸣潮 CCV 采集守护 + 看板")
    parser.add_argument("--port", type=int, default=int(cfg.get("dashboard_port", 8088)))
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.StreamHandler(),
            TimedRotatingFileHandler(BASE_DIR / "server.log", when="midnight", backupCount=30, encoding="utf-8"),
        ],
    )

    env = load_env()
    twitch_creds = {
        "client_id": env.get("TWITCH_CLIENT_ID", ""),
        "client_secret": env.get("TWITCH_CLIENT_SECRET", ""),
    }
    if not twitch_creds["client_id"] or not twitch_creds["client_secret"]:
        logging.warning("未配置 Twitch 凭据，Twitch 任务将无法启动")
    manager = TaskManager(twitch_creds)
    Handler.manager = manager
    Handler.page = build_index_html(cfg.get("youtube_tool_url", ""))
    server = ThreadingHTTPServer(("127.0.0.1", args.port), Handler)
    logging.info("控制台已启动: http://127.0.0.1:%s/", args.port)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        manager.shutdown()
        for thread in REPORT_THREADS:
            thread.join(timeout=300)


if __name__ == "__main__":
    main()
