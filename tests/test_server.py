import json
import tempfile
import threading
import time
import unittest
from datetime import datetime as real_datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import server


class FixedDatetime(real_datetime):
    @classmethod
    def now(cls, tz=None):
        return cls(2026, 10, 30, 12, 34, 56, tzinfo=tz)


def make_manager(root: Path) -> server.TaskManager:
    manager = server.TaskManager.__new__(server.TaskManager)
    manager.tasks_path = root / "tasks.json"
    manager.tasks = {}
    manager.threads = {}
    manager.lock = threading.Lock()
    manager.twitch_creds = {"client_id": "test", "client_secret": "test"}
    manager.start_thread = lambda task_id: None
    return manager


class ExcelExportTests(unittest.TestCase):
    def test_export_contains_complete_sample_attributes_and_safe_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manager = make_manager(root)
            task_dir = root / "data" / "task-1"
            streams_dir = task_dir / "streams"
            streams_dir.mkdir(parents=True)
            (streams_dir / "tenz_987.csv").write_text(
                "ts_utc,viewer_count,title,language,game_name,tags\n"
                "2026-10-30T00:00:00+00:00,120,First title,en,VALORANT,[]\n"
                "2026-10-30T00:00:30Z,135,=DANGEROUS(),en,VALORANT,[]\n"
                "legacy-time,140,Last title,en,VALORANT,[]\n",
                encoding="utf-8-sig",
            )
            (task_dir / "streams_summary.csv").write_text(
                "stream_id,channel,display_name,title\n987,tenz,,Summary title\n",
                encoding="utf-8-sig",
            )

            with patch.object(server, "BASE_DIR", root):
                payload = manager.export_xlsx("task-1")

            workbook = load_workbook(BytesIO(payload))
            sheet = workbook["采样数据"]
            self.assertEqual(
                ["采样时间（北京时间）", "直播链接", "主播名称", "直播标题", "CCV", "直播场次ID"],
                [cell.value for cell in sheet[1]],
            )
            self.assertEqual(real_datetime(2026, 10, 30, 8, 0), sheet["A2"].value)
            self.assertEqual("yyyy-mm-dd hh:mm:ss", sheet["A2"].number_format)
            self.assertEqual("https://www.twitch.tv/tenz", sheet["B2"].value)
            self.assertEqual("https://www.twitch.tv/tenz", sheet["B2"].hyperlink.target)
            self.assertEqual("tenz", sheet["C2"].value)
            self.assertEqual("First title", sheet["D2"].value)
            self.assertEqual(120, sheet["E2"].value)
            self.assertEqual("987", sheet["F2"].value)
            self.assertEqual("'=DANGEROUS()", sheet["D3"].value)
            self.assertNotEqual("f", sheet["D3"].data_type)
            self.assertEqual("legacy-time", sheet["A4"].value)
            workbook.close()

    def test_export_creates_one_native_chart_with_a_series_per_stream(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manager = make_manager(root)
            task_dir = root / "data" / "task-1"
            streams_dir = task_dir / "streams"
            streams_dir.mkdir(parents=True)
            (streams_dir / "alpha_111.csv").write_text(
                "ts_utc,viewer_count,title\n"
                "2026-10-30T00:00:00Z,10,Alpha live\n"
                "2026-10-30T00:00:30Z,20,Alpha live\n",
                encoding="utf-8-sig",
            )
            (streams_dir / "beta_222.csv").write_text(
                "ts_utc,viewer_count,title\n"
                "2026-10-30T00:00:00Z,30,Beta live\n"
                "2026-10-30T00:00:30Z,40,Beta live\n",
                encoding="utf-8-sig",
            )
            (task_dir / "streams_summary.csv").write_text(
                "stream_id,channel,display_name,title\n"
                "111,alpha,Alpha,Alpha live\n"
                "222,beta,Beta,Beta live\n",
                encoding="utf-8-sig",
            )

            with patch.object(server, "BASE_DIR", root):
                payload = manager.export_xlsx("task-1")

            workbook = load_workbook(BytesIO(payload))
            chart_sheet = workbook["CCV趋势图"]
            self.assertEqual(1, len(chart_sheet._charts))
            chart = chart_sheet._charts[0]
            self.assertEqual(2, len(chart.series))
            self.assertEqual("'采样数据'!$E$2:$E$3", chart.series[0].val.numRef.f)
            self.assertEqual("'采样数据'!$E$4:$E$5", chart.series[1].val.numRef.f)
            self.assertEqual("'采样数据'!$A$2:$A$3", chart.series[0].cat.numRef.f)
            self.assertEqual("采样时间（北京时间）", chart.x_axis.title.tx.rich.p[0].r[0].t)
            self.assertEqual("CCV", chart.y_axis.title.tx.rich.p[0].r[0].t)
            workbook.close()

    def test_export_with_only_empty_stream_files_has_no_empty_chart(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manager = make_manager(root)
            task_dir = root / "data" / "task-1"
            streams_dir = task_dir / "streams"
            streams_dir.mkdir(parents=True)
            (streams_dir / "tenz_987.csv").write_text(
                "ts_utc,viewer_count,title\n",
                encoding="utf-8-sig",
            )
            (task_dir / "streams_summary.csv").write_text(
                "stream_id,channel,display_name,title\n987,tenz,TenZ,Summary title\n",
                encoding="utf-8-sig",
            )

            with patch.object(server, "BASE_DIR", root):
                payload = manager.export_xlsx("task-1")

            workbook = load_workbook(BytesIO(payload))
            self.assertEqual(["采样数据", "CCV趋势图"], workbook.sheetnames)
            self.assertEqual(0, len(workbook["CCV趋势图"]._charts))
            workbook.close()


class TaskIdTests(unittest.TestCase):
    def test_create_same_target_in_same_second_keeps_both_tasks(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            payload = {"platform": "twitch", "mode": "game", "target": "Wuthering Waves"}

            with patch.object(server, "datetime", FixedDatetime):
                first = manager.create(payload)
                second = manager.create(payload)

            self.assertNotEqual(first["id"], second["id"])
            self.assertEqual(2, len(manager.tasks))
            saved = json.loads(manager.tasks_path.read_text(encoding="utf-8"))
            self.assertEqual(2, len(saved["tasks"]))

    def test_create_rejects_youtube_even_when_legacy_key_exists(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            manager.youtube_key = "legacy-key"

            with self.assertRaisesRegex(ValueError, "只支持 Twitch"):
                manager.create(
                    {
                        "platform": "youtube",
                        "mode": "video",
                        "target": "https://www.youtube.com/watch?v=abc123",
                    }
                )

            self.assertEqual({}, manager.tasks)
            self.assertFalse(manager.tasks_path.exists())


class IntervalConfigurationTests(unittest.TestCase):
    def test_channel_task_defaults_both_intervals_to_thirty_seconds(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            with patch.object(server, "datetime", FixedDatetime):
                task = manager.create(
                    {"platform": "twitch", "mode": "channel", "target": "https://twitch.tv/tenz"}
                )

            self.assertEqual(30, task.get("watch_interval_seconds"))
            self.assertEqual(30, task.get("live_interval_seconds"))
            self.assertEqual(30, task["poll_interval_seconds"])

    def test_channel_task_accepts_independent_intervals_and_clamps_minimum(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            with patch.object(server, "datetime", FixedDatetime):
                task = manager.create(
                    {
                        "platform": "twitch",
                        "mode": "channel",
                        "target": "tenz",
                        "watch_interval_seconds": 1,
                        "live_interval_seconds": 75,
                    }
                )

            self.assertEqual(5, task.get("watch_interval_seconds"))
            self.assertEqual(75, task.get("live_interval_seconds"))
            self.assertEqual(75, task["poll_interval_seconds"])

    def test_api_config_falls_back_to_legacy_poll_interval(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            task = {
                "id": "twitch_tenz_123456",
                "platform": "twitch",
                "mode": "channel",
                "channels": ["tenz"],
                "poll_interval_seconds": 45,
            }

            cfg = manager._api_cfg(task, "data/twitch_tenz_123456")

            self.assertEqual(45, cfg.get("watch_interval_seconds"))
            self.assertEqual(45, cfg.get("live_interval_seconds"))
            self.assertEqual(45, cfg["poll_interval_seconds"])

    def test_dual_channel_uses_live_interval_for_both_api_phases(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            task = {
                "id": "twitch_tenz_123456",
                "platform": "twitch",
                "mode": "channel",
                "channels": ["tenz"],
                "dual": True,
                "watch_interval_seconds": 300,
                "live_interval_seconds": 30,
                "poll_interval_seconds": 30,
            }

            cfg = manager._api_cfg(task, "data/twitch_tenz_123456")

            self.assertEqual(30, cfg["watch_interval_seconds"])
            self.assertEqual(30, cfg["live_interval_seconds"])


class ExternalToolLinkTests(unittest.TestCase):
    def normalizer(self):
        fn = getattr(server, "normalize_external_tool_url", None)
        if fn is None:
            self.fail("normalize_external_tool_url is not implemented")
        return fn

    def renderer(self):
        fn = getattr(server, "build_index_html", None)
        if fn is None:
            self.fail("build_index_html is not implemented")
        return fn

    def test_normalizer_accepts_http_and_https_urls(self):
        normalize = self.normalizer()
        self.assertEqual("https://tools.example/youtube", normalize("https://tools.example/youtube"))
        self.assertEqual("http://127.0.0.1:9000/", normalize(" http://127.0.0.1:9000/ "))

    def test_normalizer_rejects_non_external_or_credential_bearing_urls(self):
        normalize = self.normalizer()
        unsafe = [
            "",
            "/youtube",
            "javascript:alert(1)",
            "file:///C:/tool.html",
            "https://user:secret@tools.example/",
            "https://[",
        ]
        self.assertEqual(["", "", "", "", "", ""], [normalize(value) for value in unsafe])

    def test_renderer_enables_configured_external_button_safely(self):
        html = self.renderer()("https://tools.example/youtube?a=1&b=2")
        self.assertIn('href="https://tools.example/youtube?a=1&amp;b=2"', html)
        self.assertIn('target="_blank"', html)
        self.assertIn('rel="noopener noreferrer"', html)
        self.assertIn("打开 YouTube 采集工具", html)

    def test_renderer_disables_button_when_url_is_invalid(self):
        html = self.renderer()("javascript:alert(1)")
        self.assertNotIn("javascript:alert(1)", html)
        self.assertIn("YouTube 采集工具（待配置）", html)
        self.assertIn("aria-disabled=\"true\"", html)


class DashboardIntervalTests(unittest.TestCase):
    def test_dashboard_renders_two_interval_controls_and_persistent_watch_copy(self):
        page = server.build_index_html("")
        self.assertIn('id="watchInterval"', page)
        self.assertIn('id="liveInterval"', page)
        self.assertEqual(2, page.count('value="30"'))
        self.assertNotIn('id="watch"', page)
        self.assertNotIn("下播后自动停止", page)
        self.assertIn("未开播检查间隔", page)
        self.assertIn("开播采样间隔", page)

    def test_dashboard_payload_sends_both_intervals_and_uses_heartbeat_phase(self):
        page = server.build_index_html("")
        self.assertIn("watch_interval_seconds:", page)
        self.assertIn("live_interval_seconds:", page)
        self.assertIn('hb.phase==="waiting"', page)
        self.assertNotIn("auto_stop_on_offline:", page)


class FakeWorker:
    def __init__(self, alive_after_join: bool):
        self.alive_after_join = alive_after_join

    def join(self, timeout=None):
        return None

    def is_alive(self):
        return self.alive_after_join


class DeleteTests(unittest.TestCase):
    def test_delete_timeout_keeps_task_and_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manager = make_manager(root)
            manager.tasks["task-1"] = {"id": "task-1", "status": "stopping"}
            manager.threads["task-1"] = FakeWorker(alive_after_join=True)
            data_dir = root / "data" / "task-1"
            data_dir.mkdir(parents=True)
            marker = data_dir / "samples.csv"
            marker.write_text("sample", encoding="utf-8")

            with patch.object(server, "BASE_DIR", root), self.assertLogs(level="ERROR"):
                deleted = manager._delete_after_stop("task-1", purge=True)

            self.assertFalse(deleted)
            self.assertIn("task-1", manager.tasks)
            self.assertEqual("delete_failed", manager.tasks["task-1"]["status"])
            self.assertTrue(marker.exists())

    def test_delete_after_worker_exit_removes_task_and_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manager = make_manager(root)
            manager.tasks["task-1"] = {"id": "task-1", "status": "stopping"}
            manager.threads["task-1"] = FakeWorker(alive_after_join=False)
            data_dir = root / "data" / "task-1"
            data_dir.mkdir(parents=True)
            (data_dir / "samples.csv").write_text("sample", encoding="utf-8")

            with patch.object(server, "BASE_DIR", root):
                deleted = manager._delete_after_stop("task-1", purge=True)

            self.assertTrue(deleted)
            self.assertNotIn("task-1", manager.tasks)
            self.assertFalse(data_dir.exists())

    def test_retry_delete_does_not_bypass_a_still_alive_worker(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manager = make_manager(root)
            manager.tasks["task-1"] = {"id": "task-1", "status": "delete_failed"}
            manager.threads["task-1"] = FakeWorker(alive_after_join=True)
            data_dir = root / "data" / "task-1"
            data_dir.mkdir(parents=True)
            marker = data_dir / "samples.csv"
            marker.write_text("sample", encoding="utf-8")

            with patch.object(server, "BASE_DIR", root), self.assertLogs(level="ERROR"):
                result = manager.delete("task-1", purge=True)
                deadline = time.monotonic() + 2
                while manager.tasks["task-1"]["status"] != "delete_failed" and time.monotonic() < deadline:
                    time.sleep(0.01)

            self.assertTrue(result["pending"])
            self.assertIn("task-1", manager.tasks)
            self.assertTrue(marker.exists())


class RecoveryManager(server.TaskManager):
    observed_start_states = []

    def start_thread(self, task_id: str) -> None:
        self.observed_start_states.append(self.tasks[task_id]["status"])
        self.tasks[task_id]["status"] = "running"
        self.save()


class ImmediatelyStoppedRecoveryManager(RecoveryManager):
    def start_thread(self, task_id: str) -> None:
        super().start_thread(task_id)
        self.tasks[task_id]["status"] = "stopped"
        self.save()


class RecoveryTests(unittest.TestCase):
    def test_load_marks_persisted_running_task_recovering_before_restart(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data_dir = root / "data"
            data_dir.mkdir()
            tasks_path = data_dir / "tasks.json"
            tasks_path.write_text(
                json.dumps({"tasks": [{"id": "task-1", "platform": "twitch", "status": "running"}]}),
                encoding="utf-8",
            )
            RecoveryManager.observed_start_states = []

            with patch.object(server, "BASE_DIR", root):
                manager = RecoveryManager({"client_id": "test", "client_secret": "test"})

            self.assertEqual(["recovering"], manager.observed_start_states)
            self.assertEqual("running", manager.tasks["task-1"]["status"])
            saved = json.loads(tasks_path.read_text(encoding="utf-8"))
            self.assertEqual("running", saved["tasks"][0]["status"])

    def test_start_thread_sets_running_before_worker_observes_task(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            manager.tasks["task-1"] = {"id": "task-1", "status": "recovering"}
            observed = []

            def worker(task_id):
                observed.append(manager.tasks[task_id]["status"])

            manager._worker = worker
            manager.start_thread = server.TaskManager.start_thread.__get__(manager)
            manager.start_thread("task-1")
            manager.threads["task-1"].join(timeout=2)

            self.assertEqual(["running"], observed)

    def test_load_does_not_overwrite_a_restored_worker_terminal_status(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data_dir = root / "data"
            data_dir.mkdir()
            (data_dir / "tasks.json").write_text(
                json.dumps({"tasks": [{"id": "task-1", "platform": "twitch", "status": "running"}]}),
                encoding="utf-8",
            )

            with patch.object(server, "BASE_DIR", root):
                manager = ImmediatelyStoppedRecoveryManager({"client_id": "test", "client_secret": "test"})

            self.assertEqual("stopped", manager.tasks["task-1"]["status"])

    def test_load_marks_legacy_youtube_task_unsupported_without_restart(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data_dir = root / "data"
            data_dir.mkdir()
            tasks_path = data_dir / "tasks.json"
            tasks_path.write_text(
                json.dumps(
                    {
                        "tasks": [
                            {
                                "id": "yt-old",
                                "platform": "youtube",
                                "mode": "video",
                                "status": "running",
                            }
                        ]
                    }
                ),
                encoding="utf-8",
            )
            RecoveryManager.observed_start_states = []

            with patch.object(server, "BASE_DIR", root):
                manager = RecoveryManager({"client_id": "test", "client_secret": "test"})

            self.assertEqual([], manager.observed_start_states)
            self.assertEqual("unsupported", manager.tasks["yt-old"]["status"])
            saved = json.loads(tasks_path.read_text(encoding="utf-8"))
            self.assertEqual("unsupported", saved["tasks"][0]["status"])

    def test_unsupported_legacy_task_cannot_be_resumed(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager = make_manager(Path(tmp))
            manager.tasks["yt-old"] = {
                "id": "yt-old",
                "platform": "youtube",
                "status": "unsupported",
            }

            with self.assertRaisesRegex(ValueError, "不支持恢复"):
                manager.control("yt-old", "resume")

            self.assertEqual("unsupported", manager.tasks["yt-old"]["status"])


if __name__ == "__main__":
    unittest.main()
